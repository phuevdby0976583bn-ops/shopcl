# -*- coding: utf-8 -*-
import os
import re
import threading
import openpyxl
from typing import List, Optional

from ui.colors import GREEN, RED, WHITE, YELLOW, EKL, LINE
from core.settings_manager import load_settings
from core.path_manager import is_termux_or_mobile, get_base_storage_dir

file_lock = threading.Lock()
_numbers_cache: Optional[List[str]] = None

MAX_FORGET_NUMBERS = 10000
MAX_CONFIRM_ENTRIES = 5000
CONFIRM_FILE = "Confirm_List.txt"
NUMBER_LIST_FILE = "Number_List.txt"


# Locate the active Number_List.txt path (checking mobile storage folder if on Termux)
def get_number_list_filepath() -> str:
    if is_termux_or_mobile():
        sdcard_num_file = os.path.join(get_base_storage_dir(), NUMBER_LIST_FILE)
        if os.path.exists(sdcard_num_file):
            return sdcard_num_file
    if os.path.exists(NUMBER_LIST_FILE):
        return NUMBER_LIST_FILE
    if os.path.exists("numbers.txt"):
        return "numbers.txt"
    return NUMBER_LIST_FILE


# Extract phone numbers from the best-matching column in an Excel file
def extract_from_excel(filename):
    try:
        wb = openpyxl.load_workbook(filename, data_only=True)
        sheet = wb.active
        target_col = None
        max_matches = 0

        # Scan first 20 rows to find the column with the most phone-like values
        for col in range(1, sheet.max_column + 1):
            matches = 0
            for row in range(2, min(22, sheet.max_row + 1)):
                val = sheet.cell(row=row, column=col).value
                if val:
                    cleaned = re.sub(r'[\s\-\(\)\+]', '', str(val).strip())
                    if cleaned.isdigit() and 7 <= len(cleaned) <= 15:
                        matches += 1
            if matches > max_matches:
                max_matches = matches
                target_col = col

        if target_col is None:
            return None, "No phone number column found."

        numbers = []
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=target_col, max_col=target_col, values_only=True):
            val = row[0]
            if val:
                cleaned = re.sub(r'[\s\-\(\)\+]', '', str(val).strip())
                if cleaned.isdigit() and 7 <= len(cleaned) <= 15:
                    numbers.append(cleaned)

        if not numbers:
            return None, "No phone numbers found in the detected column."

        return numbers, None
    except Exception as e:
        return None, str(e)


# Load phone numbers from target file (thread-safe)
def load_numbers(filepath: Optional[str] = None) -> List[str]:
    global _numbers_cache
    if filepath is None:
        filepath = get_number_list_filepath()
    with file_lock:
        if not os.path.exists(filepath):
            return []
        try:
            with open(filepath, "r", encoding="utf-8-sig", errors="ignore") as f:
                lines = [line.strip() for line in f if line.strip()]
            _numbers_cache = list(lines)
            return list(lines)
        except Exception:
            return []


# Get next number from cache or file (FIFO)
def get_next_number(filepath: Optional[str] = None) -> Optional[str]:
    global _numbers_cache
    if filepath is None:
        filepath = get_number_list_filepath()
    with file_lock:
        if _numbers_cache is None:
            if not os.path.exists(filepath):
                return None
            try:
                with open(filepath, "r", encoding="utf-8-sig", errors="ignore") as f:
                    _numbers_cache = [line.strip() for line in f if line.strip()]
            except Exception:
                return None
        if not _numbers_cache:
            return None
        return _numbers_cache.pop(0)


# Save phone numbers to target file (thread-safe)
def save_numbers(numbers: List[str], filepath: Optional[str] = None) -> None:
    global _numbers_cache
    if filepath is None:
        filepath = get_number_list_filepath()
    with file_lock:
        _numbers_cache = list(numbers)
        try:
            with open(filepath, "w", encoding="utf-8") as f:
                for n in numbers:
                    f.write(str(n).strip() + "\n")
        except Exception:
            pass


# Remove a processed number or email from active file (thread-safe)
def remove_number(number: str, filepath: Optional[str] = None) -> None:
    global _numbers_cache
    target = str(number).strip()
    if "@" in target:
        try:
            from core.email_service import remove_hotmail
            remove_hotmail(target)
        except Exception:
            pass
        return

    if filepath is None:
        filepath = get_number_list_filepath()

    with file_lock:
        if _numbers_cache is not None:
            _numbers_cache = [n for n in _numbers_cache if n != target]

        if not os.path.exists(filepath):
            return

        try:
            with open(filepath, "r", encoding="utf-8-sig", errors="ignore") as f:
                lines = [line.strip() for line in f if line.strip()]
        except Exception:
            return

        if target in lines:
            lines = [n for n in lines if n != target]

        try:
            with open(filepath, "w", encoding="utf-8") as f:
                for n in lines:
                    f.write(n + "\n")
        except Exception:
            pass


# Prompt user to input a custom file path directly
def _prompt_custom_file_path():
    while True:
        try:
            print(f"{LINE}")
            print(f" {YELLOW}[!] No input file detected automatically in target folder.")
            custom_path = input(f" {GREEN}[{RED}?{GREEN}] Enter full path to file (.txt or .xlsx) {EKL} ").strip()
            if not custom_path:
                return None

            custom_path = custom_path.strip('\"\'')

            if not os.path.exists(custom_path):
                print(f"{RED} File not found: {custom_path}")
                retry = input(f" {WHITE}Try again? (Y/n) {EKL} ").strip().lower()
                if retry in ("n", "no"):
                    return None
                continue

            if custom_path.endswith(".xlsx"):
                nums, err = extract_from_excel(custom_path)
                if nums:
                    if len(nums) > MAX_FORGET_NUMBERS:
                        print(f"{RED} Too many numbers! Maximum {MAX_FORGET_NUMBERS} allowed.")
                        return None
                    save_numbers(nums)
                    print(f" {GREEN}[{RED}?{GREEN}] Extracted {len(nums)} numbers from {custom_path}")
                    return nums
                else:
                    print(f"{RED} Failed to extract numbers from Excel: {err}")
                    return None
            else:
                numbers = load_numbers(custom_path)
                if not numbers:
                    print(f"{WHITE} File is empty.")
                    return None
                if len(numbers) > MAX_FORGET_NUMBERS:
                    print(f"{RED} Too many numbers! Maximum {MAX_FORGET_NUMBERS} allowed.")
                    return None
                print(f" {GREEN}[{RED}?{GREEN}] Loaded {len(numbers)} numbers from {custom_path}")
                return numbers
        except KeyboardInterrupt:
            raise
        except Exception as e:
            print(f"{RED} Error: {e}")
            return None


# Route file input based on settings and environment (Mobile / PC)
def process_file_input():
    settings = load_settings()
    file_cfg = settings.get("file_input_settings", {})
    always_txt = file_cfg.get("always_use_txt", False)
    multi_excel = file_cfg.get("use_multiple_excel_files", False)

    search_dir = get_base_storage_dir() if is_termux_or_mobile() else "."

    if always_txt:
        return _load_txt(search_dir)
    if multi_excel:
        return _load_multi_excel(search_dir)
    return _load_auto(search_dir)


# Load numbers directly from Number_List.txt
def _load_txt(search_dir="."):
    txt_path = os.path.join(search_dir, NUMBER_LIST_FILE)
    if not os.path.exists(txt_path):
        if is_termux_or_mobile():
            return _prompt_custom_file_path()
        print(f"{WHITE} '{NUMBER_LIST_FILE}' file was not found.")
        return []

    numbers = load_numbers(txt_path)
    if not numbers:
        print(f"{WHITE} '{txt_path}' file is empty.")
        if is_termux_or_mobile():
            return _prompt_custom_file_path()
        return []

    if len(numbers) > MAX_FORGET_NUMBERS:
        print(f"{RED} Too many numbers! Maximum {MAX_FORGET_NUMBERS} allowed.")
        print(f"{RED} You have {len(numbers)} numbers. Please reduce and try again.")
        return None

    print(f" {GREEN}[{RED}?{GREEN}] Selected File {EKL} {txt_path}")
    return numbers


# Load and merge numbers from all Excel files in the target directory
def _load_multi_excel(search_dir="."):
    xlsx_files = [
        os.path.join(search_dir, f) for f in os.listdir(search_dir)
        if f.endswith(".xlsx") and not f.startswith("~$")
    ] if os.path.exists(search_dir) else []

    if not xlsx_files:
        return _load_txt(search_dir)

    print(f" {GREEN}[{RED}?{GREEN}] Found {len(xlsx_files)} Excel Files in {search_dir}:")
    all_numbers = []
    for f in xlsx_files:
        print(f"{WHITE} Extracting from {EKL} {os.path.basename(f)}...")
        nums, err = extract_from_excel(f)
        if nums:
            all_numbers.extend(nums)
            print(f"{GREEN}  -> Found {len(nums)} numbers.")
        else:
            print(f"{RED}  -> Failed: {err}")

    if not all_numbers:
        print(f"{RED} No valid numbers found in any Excel files.")
        if is_termux_or_mobile():
            return _prompt_custom_file_path()
        return None

    all_numbers = list(set(all_numbers))

    if len(all_numbers) > MAX_FORGET_NUMBERS:
        print(f"{RED} Too many numbers! Maximum {MAX_FORGET_NUMBERS} allowed.")
        print(f"{RED} Total found: {len(all_numbers)}. Please reduce files and try again.")
        return None

    save_numbers(all_numbers)
    print(f"\n {GREEN}[{RED}?{GREEN}] Total Unique Numbers {EKL} {len(all_numbers)}")
    print(f" {GREEN}[{RED}?{GREEN}] Saved to '{NUMBER_LIST_FILE}'\n")
    return all_numbers


# Auto-detect input source: use Excel if available, fallback to txt or custom path
def _load_auto(search_dir="."):
    xlsx_files = [
        os.path.join(search_dir, f) for f in os.listdir(search_dir)
        if f.endswith(".xlsx") and not f.startswith("~$")
    ] if os.path.exists(search_dir) else []

    txt_path = os.path.join(search_dir, NUMBER_LIST_FILE)

    if not xlsx_files:
        if os.path.exists(txt_path):
            return _load_txt(search_dir)
        if is_termux_or_mobile():
            return _prompt_custom_file_path()
        return _load_txt(search_dir)

    filename = None
    if len(xlsx_files) == 1:
        filename = xlsx_files[0]
    else:
        print(f" {GREEN}[{RED}?{GREEN}] Found {len(xlsx_files)} Excel Files in {search_dir}:")
        for idx, f in enumerate(xlsx_files, 1):
            print(f" {GREEN}[{RED}{idx}{GREEN}] {os.path.basename(f)}")
        print(f"{LINE}")
        while True:
            try:
                choice = input(f" {GREEN}[{RED}?{GREEN}] Select File (1-{len(xlsx_files)}) {EKL} ").strip()
                if choice.isdigit():
                    idx = int(choice) - 1
                    if 0 <= idx < len(xlsx_files):
                        filename = xlsx_files[idx]
                        break
                print(f"{RED} Invalid selection!")
            except KeyboardInterrupt:
                raise
            except Exception:
                pass

    print(f" {GREEN}[{RED}?{GREEN}] Selected File {EKL} {filename}\n")
    nums, err = extract_from_excel(filename)

    if nums:
        if len(nums) > MAX_FORGET_NUMBERS:
            print(f"{RED} Too many numbers! Maximum {MAX_FORGET_NUMBERS} allowed.")
            print(f"{RED} Found {len(nums)} numbers in {filename}. Please reduce and try again.")
            return None
        save_numbers(nums)
        print(f" {GREEN}[{RED}?{GREEN}] Extracted {len(nums)} numbers from {filename}")
        print(f" {GREEN}[{RED}?{GREEN}] Saved to '{NUMBER_LIST_FILE}'\n")
        return nums
    else:
        print(f"{RED} Error: {err}")
        if is_termux_or_mobile():
            return _prompt_custom_file_path()
        return None
